"""Parse `Pokémon Locations.xlsx`.

Encounter sheets are laid out as one column *pair* per location:

    row 1   |        | Littleroot Town |        | Route 101 |
    row 2   |        | Level | Pokémon | Level | Pokémon   |
    col 1   method name at the first row of each block ("Land", "Fishing", ...)
    col 2   slot rate

Excel helpfully turned the level range "2-3" into a date, so datetimes are
turned back into day-month text.
"""

import datetime
import re

import openpyxl

from . import names

ENCOUNTER_SHEETS = {
    "Encounter Tables": "normal",
    "Temp Encounter Tables": "temp",
}
SPECIAL_TAGS = ("Static", "Gift")


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        # "2-3" was read as 2 March; put the original text back together.
        return f"{value.day}-{value.month}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rate(value):
    """Slot rates are stored as '20%' in some blocks and 0.2 in others."""
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip().rstrip("%")
        return round(float(text), 2) if text else None
    return round(float(value) * 100, 2)


def _slot(raw_name, method, rate, levels):
    tag = None
    for candidate in SPECIAL_TAGS:
        if raw_name.startswith(candidate + " "):
            tag, raw_name = candidate, raw_name[len(candidate) + 1:]
            break

    form = names.COSMETIC_FORMS.get(raw_name)
    expanded = names.expand_encounter_name(raw_name)
    return {
        "method": method,
        "rate": rate,
        "levels": levels,
        "name": names.display(expanded),
        "key": names.key(expanded),
        "form": form[1] if form else None,
        "tag": tag,
    }


def _read_encounter_sheet(worksheet):
    """-> {location name: [slot, ...]}"""
    columns = {}
    for column in range(1, worksheet.max_column + 1):
        location = _cell(worksheet.cell(1, column).value)
        if location:
            columns[column] = location

    result = {location: [] for location in columns.values()}
    method = None
    for row in range(3, worksheet.max_row + 1):
        label = _cell(worksheet.cell(row, 1).value)
        if label:
            method = label
        rate = _rate(worksheet.cell(row, 2).value)
        if method is None:
            continue
        for column, location in columns.items():
            levels = _cell(worksheet.cell(row, column).value)
            raw_name = _cell(worksheet.cell(row, column + 1).value)
            if raw_name:
                result[location].append(_slot(raw_name, method, rate, levels))
    return result


def _read_notes(worksheet):
    """Split a free-form sheet into blocks separated by blank rows."""
    blocks, current = [], []
    for row in worksheet.iter_rows(values_only=True):
        cells = [_cell(c) for c in row]
        if any(cells):
            current.append(cells)
        elif current:
            blocks.append(current)
            current = []
    if current:
        blocks.append(current)
    return blocks


def parse(path):
    workbook = openpyxl.load_workbook(path, data_only=True)

    tables = {}
    for sheet_name, table_set in ENCOUNTER_SHEETS.items():
        tables[table_set] = _read_encounter_sheet(workbook[sheet_name])

    underwater = _read_encounter_sheet(workbook["Underwater Encounter Tables"])
    for table_set in tables.values():
        for location, slots in underwater.items():
            table_set.setdefault(location, []).extend(slots)

    ordered = list(tables["normal"].keys())
    locations = []
    for name in ordered:
        normal = tables["normal"].get(name, [])
        temp = tables["temp"].get(name, [])
        locations.append({
            "name": name,
            "slots": normal,
            # Only four locations differ between the two sheets; the rest show
            # a single table.
            "tempSlots": temp if temp != normal else None,
        })

    return {
        "locations": locations,
        "gameCorner": _parse_game_corner(workbook["Game Corner"]),
        **_parse_other(workbook["Other"]),
        "unavailable": _parse_unavailable(workbook["Unavailable Pokémon"]),
    }


def _parse_game_corner(worksheet):
    blocks = _read_notes(worksheet)
    rows = [row for block in blocks for row in block]
    note = rows[0][0]
    rewards = []
    for cells in rows[1:]:
        if cells[0] and cells[0].lower() != "reward pokémon" and cells[1]:
            rewards.append({"badge": cells[0], "pokemon": _mon_list(cells[1])})
    return {"note": note, "rewards": rewards}


def _mon_list(text):
    """'Smoochum, Elekid or Magby' -> linkable entries."""
    parts = re.split(r",\s*|\s+or\s+", text.strip().rstrip("."))
    return [{"name": p.strip(), "key": names.key(p)} for p in parts if p.strip()]


def _parse_other(worksheet):
    blocks = _read_notes(worksheet)
    result = {"trades": [], "gifts": [], "fossils": [], "roaming": []}
    for block in blocks:
        heading = block[0][0]
        body = [cells for cells in block[2:] if any(cells)]
        if heading.startswith("Trades"):
            result["trades"] = [
                {"wanted": c[0], "location": c[1], "offered": c[2]} for c in body
            ]
        elif heading.startswith("Gifts"):
            result["gifts"] = [{"location": c[0], "pokemon": c[1]} for c in body]
        elif heading.startswith("Fossils"):
            result["fossils"] = [{"location": c[0], "detail": c[1]} for c in body]
        elif heading.startswith("Roaming"):
            result["roaming"] = [
                {"routes": _flatten(c[0]), "pokemon": _mon_list(_flatten(c[1]))}
                for c in body
            ]
    return result


def _flatten(text):
    return re.sub(r"\s*\n\s*", " ", text).strip()


def _parse_unavailable(worksheet):
    groups = []
    for cells in _read_notes(worksheet)[0][1:]:
        if cells[0] and cells[1]:
            groups.append({
                "generation": cells[0],
                "pokemon": [p.strip() for p in cells[1].rstrip(".").split(",")],
            })
    return groups
